"""Regression tests for the TXT > Word upload route and tagged Zuid structure."""
import base64
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
import app


# Synthetic data: structural properties from the supplied South example, no article text.
ZUID_STRUCTURE = '''<body>\r
<l_region>\r
Heerlen<EP>\r
</l_region>\r
<l_info><bold>A</bold><SU>Adres A * za 17-10 <EP></l_info>
<l_info><EP></l_info>
<l_region>Vaals<EP></l_region>
<l_info><bold>B</bold><SU>Adres B<EP></l_info>
<l_region><EP></l_region><l_info><EP></l_info>
<EP><l_region>Kerkrade<EP></l_region>
<l_info(1)><bold(1)>C</bold(1)><SU>Adres C<extra_leading><EP></l_info(1)>
<l_region>Heerlen<EP></l_region>
<l_info><bold>A</bold><SU>Adres A * za 17-10 <EP></l_info>
<l_info><bold>D</bold><SU>Adres D<extra_leading><extra_leading><EP></l_info>
</body>'''


class UploadTests(unittest.TestCase):
    def upload(self, raw: bytes, filename='agenda zuid.txt'):
        stream = BytesIO(raw)
        stream.filename = filename
        fake_request = SimpleNamespace(
            method='POST', form={'export_type': 'word'}, files={'txt_file': stream}
        )
        with patch.object(app, 'request', fake_request), patch.object(
            app, 'render_main', side_effect=lambda **kwargs: kwargs
        ):
            return app.index()

    def test_zuid_tag_variants_crlf_repeated_blocks_and_empty_sections(self):
        result = self.upload(ZUID_STRUCTURE.encode('utf-8'))
        self.assertNotIn('error', result)
        doc = Document(BytesIO(base64.b64decode(result['word_docx_base64'])))
        self.assertEqual([p.text for p in doc.paragraphs], [
            'HEERLEN', 'A Adres A * za 17-10', 'D Adres D',
            'KERKRADE', 'C Adres C', 'VAALS', 'B Adres B'
        ])
        self.assertEqual(result['word_copy_text'].count('A Adres A * za 17-10'), 1)

    def test_header_and_footer_version_are_same(self):
        template = (Path(__file__).resolve().parent.parent /
                    'templates' / 'index.html').read_text('utf-8')
        self.assertIn('Versie {{ app_version }}', template)
        self.assertEqual(app.APP_VERSION, '20261021-txt-word-3')

    def test_windows_1252_encoded_input(self):
        raw = '<body><l_region>Bunde</l_region><l_info><bold>Résumé</bold><SU>€ 5</l_info></body>'.encode('cp1252')
        result = self.upload(raw)
        self.assertNotIn('error', result)
        doc = Document(BytesIO(base64.b64decode(result['word_docx_base64'])))
        self.assertEqual(doc.paragraphs[1].text, 'Résumé € 5')

    def test_invalid_tag_gives_technical_detail_only(self):
        raw = '<body><l_region>Venlo</l_region><l_info><bold>Nieuws</l_info></body>'.encode()
        result = self.upload(raw)
        self.assertIn('ERR_TXT_FORMAT (BOLD_UNCLOSED)', result['error'])
        self.assertNotIn('Nieuws', result['error'])

    def test_bad_text_encoding_gets_distinct_error(self):
        result = self.upload(b'\x81')
        self.assertIn('ERR_TXT_ENCODING', result['error'])


if __name__ == '__main__':
    unittest.main()
