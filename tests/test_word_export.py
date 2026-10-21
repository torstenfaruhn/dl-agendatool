"""Tests voor getagde Cue Print-TXT naar een opgemaakt Word-document."""
from io import BytesIO
import unittest

from docx import Document

from app import (InvalidTaggedAgenda, build_word_copy_formats,
                 parse_tagged_agenda, transform_txt_to_docx)


SAMPLE_TXT = '''<body>
<l_region>
Bunde<EP>
</l_region>
<EP>
<l_info><bold>
Het Landschap verbeeld door Arjen van Prooijen
</bold>
<SU>De Agnespoort, Agnesplein 14, 6241 CA Bunde * do 15-10 * 20.00-22.00 uur
<EP></l_info>
<l_region>Eijsden<EP></l_region>
<l_info><bold>Rommelmarkt </bold>
Cultureel Centrum Eijsden, Breusterstraat 23, 6245 EG Eijsden * zo 18-10 *
9.00-17.00 uur<EP></l_info>
</body>'''


def read_docx(source):
    return Document(BytesIO(transform_txt_to_docx(source)))


class WordExportTests(unittest.TestCase):
    def test_example_text_matches_user_word_example(self):
        doc = read_docx(SAMPLE_TXT)
        self.assertEqual([p.text for p in doc.paragraphs], [
            "BUNDE",
            "Het Landschap verbeeld door Arjen van Prooijen De Agnespoort, Agnesplein 14, 6241 CA Bunde * do 15-10 * 20.00-22.00 uur",
            "EIJSDEN",
            "Rommelmarkt Cultureel Centrum Eijsden, Breusterstraat 23, 6245 EG Eijsden * zo 18-10 * 9.00-17.00 uur",
        ])

    def test_bold_is_real_word_formatting_not_visible_tags(self):
        doc = read_docx(SAMPLE_TXT)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertFalse(bool(doc.paragraphs[1].runs[-1].bold))
        self.assertTrue(doc.paragraphs[2].runs[0].bold)
        self.assertTrue(doc.paragraphs[3].runs[0].bold)
        self.assertFalse(bool(doc.paragraphs[3].runs[-1].bold))
        combined = " ".join(p.text for p in doc.paragraphs)
        for tag in ('<body>', '<l_region>', '<l_info>', '<bold>', '<SU>', '<EP>'):
            self.assertNotIn(tag, combined)

    def test_handles_numbered_info_empty_sections_and_extra_leading(self):
        src = '''<body><l_region>Vaals<EP></l_region>
        <l_info><EP></l_info><l_info(1)><bold>Nieuws</bold><SU>De zaal<extra_leading><EP></l_info(1)>
        <l_region><EP></l_region><l_info><EP></l_info>
        </body>'''
        self.assertEqual([p.text for p in read_docx(src).paragraphs], ["VAALS", "Nieuws De zaal"])

    def test_preserves_source_order_and_repeated_place_sections(self):
        src = '''<body>
        <l_region>Zutendaal</l_region><l_info><bold>A</bold><SU>Een</l_info>
        <l_region>Arcen</l_region><l_info><bold>B</bold><SU>Twee</l_info>
        <l_region>Zutendaal</l_region><l_info><bold>A</bold><SU>Een</l_info>
        </body>'''
        self.assertEqual([p.text for p in read_docx(src).paragraphs], [
            "ZUTENDAAL", "A Een", "ARCEN", "B Twee", "ZUTENDAAL", "A Een"
        ])

    def test_keeps_content_without_su_or_without_bold(self):
        src = '''<body><l_region>Arcen</l_region>
        <l_info><bold>Titel</bold>Locatie * 20.00 uur<EP></l_info>
        <l_info>Bericht zonder vetopmaak<EP></l_info>
        </body>'''
        doc = read_docx(src)
        self.assertEqual([p.text for p in doc.paragraphs], [
            "ARCEN", "Titel Locatie * 20.00 uur", "Bericht zonder vetopmaak"
        ])
        self.assertFalse(bool(doc.paragraphs[-1].runs[-1].bold))

    def test_date_and_punctuation_not_normalized_again(self):
        src = '''<body><l_region>Arcen</l_region><l_info>
        <bold>Start</bold><SU>9 uur * € 7,– & details
        </l_info></body>'''
        self.assertEqual(read_docx(src).paragraphs[1].text, "Start 9 uur * € 7,– & details")

    def test_copy_format_html_is_escaped(self):
        src = '''<body><l_region>Arcen</l_region>
        <l_info><bold>Concert & Muziek</bold><SU>Firma A & Co</l_info></body>'''
        plain, html = build_word_copy_formats(parse_tagged_agenda(src))
        self.assertIn('Concert & Muziek Firma A & Co', plain)
        self.assertIn('<strong>Concert &amp; Muziek</strong>', html)
        self.assertIn('Firma A &amp; Co', html)

    def test_plain_text_instead_of_tagged_agenda_is_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('Gewone tekst zonder agenda-tags')

    def test_unclosed_region_is_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('<body><l_region>Arcen<l_info><bold>Foo</bold>x</l_info></body>')

    def test_unclosed_bold_is_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('<body><l_region>Arcen</l_region><l_info><bold>Foo</l_info></body>')

    def test_event_without_place_is_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('<body><l_info><bold>Foo</bold>Text</l_info></body>')

    def test_empty_sections_only_are_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('<body><l_region>Arcen</l_region><l_info><EP></l_info></body>')

    def test_wrong_numbered_info_is_rejected(self):
        with self.assertRaises(InvalidTaggedAgenda):
            parse_tagged_agenda('<body><l_region>Arcen</l_region><l_info(1)><bold>Foo</bold>x</l_info(2)></body>')


if __name__ == "__main__":
    unittest.main()
