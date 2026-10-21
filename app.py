# app.py - XML naar Cue Print (TXT); getagde TXT naar Cue Web (DOCX)
# Generated for deployment on Render.com

import base64
from html import escape
from io import BytesIO
import secrets

from docx import Document
from docx.shared import Cm, Pt
from flask import Flask, g, render_template, request

# === Transformatielogica uit het Colab-notebook ===
# Transformatiefunctie met verbeterde datum- en tijdlogica
import re
import xml.etree.ElementTree as ET
from collections import defaultdict

MONTHS = [
    "januari","februari","maart","april","mei","juni",
    "juli","augustus","september","oktober","november","december"
]
MONTH_INDEX = {m: i+1 for i, m in enumerate(MONTHS)}

def normalize_place_name(t: str) -> str:
    if not t:
        return ''
    t = t.strip()
    repl = [
        (r"\bUrmond\s+Gemeente\s+Stein\b", "Urmond"),
        (r"\bValkenburg\s+aan\s+de\s+Geul\b", "Valkenburg"),
        (r"\bValkenburg\s+a/d\s+Geul\b", "Valkenburg"),
        (r"\bElsoo\s+Lb\b", "Elsloo"),
        (r"\bGemeente\s+Stein\b", "Stein"),
        (r"\bNederweert\s+Eind\b", "Nederweert-Eind"),
        (r"\bHout\s+Blerick\b", "Hout-Blerick"),
    ]
    for p, r in repl:
        t = re.sub(p, r, t, flags=re.IGNORECASE)
    t = re.sub(r"\s+", " ", t)
    return t.strip()

def normalize_postcodes(t: str) -> str:
    if not t:
        return t
    return re.sub(r"\b(\d{4})\s*([A-Za-z]{2})\b",
                  lambda m: f"{m.group(1)} {m.group(2).upper()}", t)

def normalize_housenumbers(t: str) -> str:
    if not t:
        return t
    return re.sub(r"\b(\d+)\s*([A-Za-z])\b",
                  lambda m: f"{int(m.group(1))}{m.group(2).lower()}", t)

def normalize_times(t: str) -> str:
    """Implementeert regel 13, zonder minuten aan te passen of dubbele .00 te maken."""
    if not t:
        return t
    # Verwijder spaties na '.' in tijd-achtige patronen: 10. 00 -> 10.00, 15. 35 -> 15.35
    t = re.sub(r"(\d{1,2})\.\s*(\d{2})", r"\1.\2", t)
    # 10-14.30 uur -> 10.00-14.30 uur (alleen als eerste deel een uur is 0-23)
    t = re.sub(r"(?<!\.)\b([01]?\d|2[0-3])-(\d{1,2}\.\d{2})\s*uur\b",
              lambda m: f"{int(m.group(1))}.00-{m.group(2)} uur", t)
    # 8-10 uur -> 8.00-10.00 uur
    t = re.sub(r"(?<!\.)\b([01]?\d|2[0-3])-(\d{1,2})\s*uur\b",
              lambda m: f"{int(m.group(1))}.00-{int(m.group(2))}.00 uur", t)
    # 9 uur -> 9.00 uur (en alleen uren 0-23, niet '35 uur' in 15.35 uur)
    t = re.sub(r"(?<!\.)\b([01]?\d|2[0-3])\s*uur\b",
              lambda m: f"{int(m.group(1))}.00 uur", t)
    # 12u -> 12.00 uur
    t = re.sub(r"(?<!\.)\b([01]?\d|2[0-3])u\b",
              lambda m: f"{int(m.group(1))}.00 uur", t)
    return t

def simplify_date_ranges(t: str) -> str:
    """Implementeert regels 11 en 12.

    - Herken datumbereiken "datum1 ... tot/t/m datum2" en maak er
      "* t/m [datum2] *" of "* t/m [weekday datum2]" van.
    - Ondersteunt zowel maandnamen als numerieke datums met weekdag.
    """
    if not t:
        return t

    # 1) Numerieke datums met weekdag en tijden, zoals:
    #    * zo 22-6 * 14.00 uur tot zo 13-7 * 18.00 uur
    weekday_pattern = r"(?:ma|di|wo|do|vr|za|zo)"

    def repl_numeric(m):
        wd2 = m.group(1)
        date2 = m.group(2)
        return f"* t/m {wd2} {date2} "

    numeric_pat = (
        rf"\*\s*{weekday_pattern}\s+\d{{1,2}}-\d{{1,2}}\s*"  # * zo 22-6
        rf"\*\s*[^*]*?(?:tot|t/m)\s*"                        # * 14.00 uur tot
        rf"({weekday_pattern})\s+(\d{{1,2}}-\d{{1,2}})"      # zo 13-7
        rf"\s*\*[^*]*"                                       # * 18.00 uur
    )

    t = re.sub(numeric_pat, repl_numeric, t, flags=re.IGNORECASE)

    # 2) Variant met maandnamen, zoals: 22 juni tot 13 juli
    months_pattern = "|".join(MONTHS)
    pat = (
        rf"(\d{{1,2}})\s+({months_pattern})(?:\s+(\d{{2,4}}))?"  # dag1 maand1 [jaar1]
        rf"\s+(?:tot|t/m)\s+"                                     # tot / t/m
        rf"(\d{{1,2}})\s+({months_pattern})(?:\s+(\d{{2,4}}))?"  # dag2 maand2 [jaar2]
    )

    def repl_month(m):
        day1 = m.group(1)
        month1 = m.group(2)
        year1 = m.group(3)
        day2 = m.group(4)
        month2 = m.group(5)
        year2 = m.group(6)

        date2_str = f"{day2} {month2}"
        if year2:
            date2_str = f"{date2_str} {year2}"
        else:
            m1_idx = MONTH_INDEX.get(month1.lower(), None)
            m2_idx = MONTH_INDEX.get(month2.lower(), None)
            if m1_idx is not None and m2_idx is not None and m2_idx < m1_idx:
                # Cross-jaar: -26 toevoegen zoals beschreven
                date2_str = f"{date2_str} -26"

        return f"* t/m {date2_str} *"

    t = re.sub(pat, repl_month, t, flags=re.IGNORECASE)
    return t

def normalize_spaces_and_punctuation(t: str) -> str:
    """Regel 14: spaties rond leestekens normaliseren, maar **zonder** tijden te breken.
    We laten de punt (.) met rust; die is belangrijk voor tijden en decimalen.
    """
    if not t:
        return t
    # Verwijder dubbele spaties
    t = re.sub(r"\s+", " ", t)
    # Behandel alleen , ; : ! ?
    t = re.sub(r"\s+([,;:!?])", r"\1", t)
    t = re.sub(r"([;:!?])(\S)", r"\1 \2", t)
    # eurobedragen
    t = re.sub(r"(€\s*\d+),\s*(\d{2}\b)", r"\1,\2", t)
    t = re.sub(r"(€\s*\d+),\s*(–|-)", r"\1,\2", t)
    return t.strip()

def strip_trailing_asterisk(t: str) -> str:
    if not t:
        return t
    t = t.rstrip()
    while t.endswith('*'):
        t = t[:-1].rstrip()
    return t

def process_su_text(t: str) -> str:
    """Volgorde van bewerkingen volgens de regels, zodat ze elkaar niet in de weg zitten.

    8. postcodes
    9. huisnummers
    10. plaatsnamen
    11/12. datumbereiken
    13. tijdstippen
    14. spaties en leestekens
    17. trailing '*'
    """
    if t is None:
        t = ""
    t = normalize_postcodes(t)
    t = normalize_housenumbers(t)
    t = normalize_place_name(t)
    t = simplify_date_ranges(t)
    t = normalize_times(t)
    t = normalize_spaces_and_punctuation(t)
    t = strip_trailing_asterisk(t)
    return t

# De bestaande XML-naar-TXT-verwerking behoudt ontdubbeling en sortering.
def collect_agenda_events(xml_input: str) -> dict[str, list[dict[str, str]]]:
    root = ET.fromstring(xml_input)
    ev_by_place = defaultdict(list)
    seen_events = set()

    for e in root.iter():
        if e.tag != "evenement":
            continue
        p = e.find("plaats")
        g = e.find("genre")
        t = e.find("tekst")
        if p is None or g is None or t is None:
            continue
        pn = normalize_place_name((p.text or "").strip())
        gn = (g.text or "").strip()
        su = process_su_text((t.text or "").strip())

        # Optie A: exact dezelfde verwerkte plaatsnaam, titel en tekst.
        event_key = (pn, gn, su)
        if event_key in seen_events:
            continue
        seen_events.add(event_key)
        ev_by_place[pn].append({"plaats": pn, "genre": gn, "su": su})

    return {
        place: ev_by_place[place]
        for place in sorted(ev_by_place, key=lambda name: name.lower())
    }


def transform_xml(xml_input: str) -> str:
    """Maak de ongewijzigde Cue Print-code (.txt) uit XML."""
    ev_by_place = collect_agenda_events(xml_input)
    lines = ["<body>"]
    for place, events in ev_by_place.items():
        lines.append(f"<l_region>{place}</l_region>")
        for ev in events:
            lines.append("<EP>")
            lines.append("<l_info>")
            lines.append(f"<bold>{ev['genre']}</bold>")
            # SU blijft bewust een opentag zonder sluittag.
            lines.append(f"<SU>{ev['su']}")
            lines.append("</l_info>")
    lines.append("</body>")
    return "\n".join(lines)


# De tweede functie gebruikt uitsluitend al bestaande, getagde TXT-invoer.
# De bron kan opeenvolgende plaatsblokken, lege blokken en afgebroken regels bevatten.
TAGGED_BLOCK_RE = re.compile(
    r"<l_region>.*?</l_region>|<l_info(?:\(\d+\))?>.*?</l_info(?:\(\d+\))?>",
    re.IGNORECASE | re.DOTALL,
)
INFO_OPEN_RE = re.compile(r"^<l_info(?P<suffix>\(\d+\))?>", re.IGNORECASE)
INFO_END_RE = re.compile(r"</l_info(?P<suffix>\(\d+\))?>$", re.IGNORECASE)
# Cues uit redactionele TXT-bestanden gebruiken zowel <bold> als <bold(1)>.
# De sluit-tag moet bij dezelfde genummerde variant horen.
BOLD_BLOCK_RE = re.compile(
    r"<bold(?P<open_suffix>\(\d+\))?>(?P<title>.*?)"
    r"</bold(?P<close_suffix>\(\d+\))?>",
    re.IGNORECASE | re.DOTALL,
)
KNOWN_TAG_RE = re.compile(
    r"</?(?:body|l_region|l_info(?:\(\d+\))?|bold(?:\(\d+\))?)>|"
    r"<(?:EP|SU|extra_leading)>",
    re.IGNORECASE,
)
SUSPECT_TAG_RE = re.compile(r"</?[A-Za-z][^<>]{0,80}>")


class InvalidTaggedAgenda(ValueError):
    """De TXT bevat geen herkenbare of complete agenda-tags."""


def _remove_txt_tags(value: str) -> str:
    without_tags = KNOWN_TAG_RE.sub("", value)
    if SUSPECT_TAG_RE.search(without_tags):
        raise InvalidTaggedAgenda("onbekende agenda-tag")
    return re.sub(r"\s+", " ", without_tags).strip()


def parse_tagged_agenda(txt_input: str) -> list[dict]:
    """Lees getagd Cue Print-TXT en verwijder exacte dubbele agendaberichten.

    De source-tags zijn opmaakmarkeringen, geen XML: <SU> en <EP> zijn
    opzettelijk niet gesloten. Herhaalde plaatsblokken worden samengevoegd.
    """
    if not txt_input or "\x00" in txt_input:
        raise InvalidTaggedAgenda("leeg of ongeldig bestand")

    blocks = list(TAGGED_BLOCK_RE.finditer(txt_input))
    if not blocks:
        raise InvalidTaggedAgenda("geen agendablokken")

    # Ontbrekende eindtags mogen niet onopgemerkt berichten laten wegvallen.
    region_count = len(re.findall(r"<l_region>", txt_input, re.IGNORECASE))
    info_count = len(re.findall(r"<l_info(?:\(\d+\))?>", txt_input, re.IGNORECASE))
    if (region_count != sum(b.group().lower().startswith("<l_region>") for b in blocks)
            or info_count != sum(b.group().lower().startswith("<l_info") for b in blocks)):
        raise InvalidTaggedAgenda("onvolledige agenda-blokken")

    sections = []
    active_section = None
    last_end = 0
    for block in blocks:
        # Buiten een plaats-/evenementblok mogen uitsluitend controle-tags staan.
        if _remove_txt_tags(txt_input[last_end:block.start()]):
            raise InvalidTaggedAgenda("tekst buiten agendablokken")
        source = block.group()
        if source.lower().startswith("<l_region>"):
            place = _remove_txt_tags(source[len("<l_region>"):-len("</l_region>")])
            active_section = {"place": place, "events": []} if place else None
            if active_section:
                sections.append(active_section)
        else:
            start_tag = INFO_OPEN_RE.match(source)
            end_tag = INFO_END_RE.search(source)
            if not start_tag or not end_tag or start_tag.group("suffix") != end_tag.group("suffix"):
                raise InvalidTaggedAgenda("onjuist genummerd agendablok")
            content = source[start_tag.end():end_tag.start()]
            open_bold = len(re.findall(r"<bold(?:\(\d+\))?>", content, re.IGNORECASE))
            close_bold = len(re.findall(r"</bold(?:\(\d+\))?>", content, re.IGNORECASE))
            bold_blocks = list(BOLD_BLOCK_RE.finditer(content))
            if open_bold != len(bold_blocks) or close_bold != len(bold_blocks):
                raise InvalidTaggedAgenda("onvolledige vetopmaak")

            segments = []
            events_in_block = []
            cursor = 0
            has_title = False
            for bold in bold_blocks:
                if bold.group("open_suffix") != bold.group("close_suffix"):
                    raise InvalidTaggedAgenda("onjuist genummerde vetopmaak")
                leading = _remove_txt_tags(content[cursor:bold.start()])
                if leading:
                    segments.append({"text": leading, "bold": False})
                # Een nieuw vetgedrukt evenement binnen hetzelfde <l_info>-blok
                # begint een nieuw bericht, ook als een <EP> ontbreekt.
                if has_title and segments:
                    events_in_block.append(segments)
                    segments = []
                title = _remove_txt_tags(bold.group("title"))
                if title:
                    segments.append({"text": title, "bold": True})
                has_title = True
                cursor = bold.end()
            following = _remove_txt_tags(content[cursor:])
            if following:
                segments.append({"text": following, "bold": False})
            if segments:
                events_in_block.append(segments)
            if events_in_block:
                if active_section is None:
                    raise InvalidTaggedAgenda("agendabericht zonder plaatsnaam")
                active_section["events"].extend(events_in_block)
        last_end = block.end()

    if _remove_txt_tags(txt_input[last_end:]):
        raise InvalidTaggedAgenda("tekst buiten agendablokken")
    # Lege placeholderblokken komen voor in het referentiebestand.
    valid_sections = [section for section in sections if section["events"]]
    if not valid_sections:
        raise InvalidTaggedAgenda("geen agendaberichten")
    return deduplicate_tagged_agenda(valid_sections)


def deduplicate_tagged_agenda(sections: list[dict]) -> list[dict]:
    """Optie A: bewaar per plaats, titel/opmaak en volledige tekst één bericht.

    De parser heeft witruimte en tags al verwerkt. Zo geldt <bold(1)> als
    dezelfde titelmarkering als <bold>. Verschillende titels of omschrijvingen
    blijven behouden. Plaatsnamen worden alfabetisch gezet, binnen een
    plaats blijft de oorspronkelijke volgorde van de eerste keer bestaan.
    """
    grouped = {}
    seen_events = set()
    for section in sections:
        place = section["place"]
        place_key = place.casefold()
        if place_key not in grouped:
            grouped[place_key] = {"place": place, "events": []}
        for segments in section["events"]:
            # Vergelijk de gehele titel en gehele overige berichttekst.
            # Afwijkingen in titel of informatie mogen niet stil verloren gaan.
            title = " ".join(part["text"] for part in segments if part["bold"])
            details = " ".join(part["text"] for part in segments if not part["bold"])
            event_key = (place_key, title, details)
            if event_key in seen_events:
                continue
            seen_events.add(event_key)
            grouped[place_key]["events"].append(segments)
    return [grouped[place] for place in sorted(grouped, key=str.casefold)]


def build_word_docx(sections: list[dict]) -> bytes:
    """Zet TXT-segmenten om naar echte Word-opmaak, niet naar zichtbare tags."""
    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    for section in sections:
        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.add_run(section["place"].upper()).bold = True
        for segments in section["events"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            for i, segment in enumerate(segments):
                if i:
                    paragraph.add_run(" ")
                paragraph.add_run(segment["text"]).bold = segment["bold"]

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def transform_txt_to_docx(txt_input: str) -> bytes:
    return build_word_docx(parse_tagged_agenda(txt_input))


def build_word_copy_formats(sections: list[dict]) -> tuple[str, str]:
    """Maak veilige HTML- en tekstvarianten voor gekopieerde Word-inhoud."""
    plain = []
    html = []
    for section in sections:
        plain.append(section["place"].upper())
        html.append(f"<p><strong>{escape(section['place'].upper())}</strong></p>")
        for segments in section["events"]:
            plain.append(" ".join(segment["text"] for segment in segments))
            html.append("<p>" + " ".join(
                f"<strong>{escape(segment['text'])}</strong>"
                if segment["bold"] else escape(segment["text"])
                for segment in segments
            ) + "</p>")
    return "\n\n".join(plain), "".join(html)


DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALLOWED_EXPORT_TYPES = {"agenda", "word"}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES


@app.before_request
def set_csp_nonce():
    g.csp_nonce = secrets.token_urlsafe(16)


@app.after_request
def add_security_headers(response):
    nonce = getattr(g, "csp_nonce", "")
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; "
        "img-src 'self' data:; "
        "style-src 'self'; "
        f"script-src 'self' 'nonce-{nonce}'; "
        "base-uri 'self'; "
        "form-action 'self'; "
        "frame-ancestors 'none'"
    )
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    return response


def render_main(**kwargs):
    defaults = {
        "output_text": None,
        "word_docx_base64": None,
        "word_copy_text": None,
        "word_copy_html": None,
        "export_type": "agenda",
        "error": None,
        "csp_nonce": getattr(g, "csp_nonce", ""),
    }
    defaults.update(kwargs)
    return render_template("index.html", **defaults)


@app.errorhandler(413)
def too_large(_error):
    return render_main(
        error="Het bestand is te groot (maximaal 5 MB). Foutcode: ERR_FILE_SIZE."
    ), 413


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "GET":
        return render_main()

    export_type = request.form.get("export_type", "agenda")
    if export_type not in ALLOWED_EXPORT_TYPES:
        return render_main(error="Onbekend exporttype. Foutcode: ERR_EXPORT_TYPE.")

    expected_ext = ".xml" if export_type == "agenda" else ".txt"
    input_key = "xml_file" if export_type == "agenda" else "txt_file"
    upload = request.files.get(input_key)
    if upload is None or upload.filename == "":
        return render_main(export_type=export_type, error=f"Geen {expected_ext}-bestand geselecteerd.")
    if not upload.filename.lower().endswith(expected_ext):
        return render_main(export_type=export_type, error=f"Kies een bestand met de extensie {expected_ext}.")

    try:
        raw = upload.read()
        if export_type == "agenda":
            # Behoud de oorspronkelijke XML-decoder en verwerkingsregels.
            try:
                xml_input = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                xml_input = raw.decode("latin-1", errors="ignore")
            return render_main(export_type=export_type, output_text=transform_xml(xml_input))

        # Voor de getagde TXT ook Windows-1252 (oudere redactionele bestanden).
        try:
            txt_input = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt_input = raw.decode("cp1252")
        sections = parse_tagged_agenda(txt_input)
        docx_bytes = build_word_docx(sections)
        copy_text, copy_html = build_word_copy_formats(sections)
        return render_main(
            export_type=export_type,
            word_docx_base64=base64.b64encode(docx_bytes).decode("ascii"),
            word_copy_text=copy_text,
            word_copy_html=copy_html,
        )
    except ET.ParseError:
        return render_main(export_type=export_type, error="Het XML-bestand is ongeldig. Foutcode: ERR_XML_FORMAT.")
    except (UnicodeDecodeError, InvalidTaggedAgenda):
        return render_main(export_type=export_type, error="Het TXT-bestand bevat ongeldige of onvolledige agendatags. Foutcode: ERR_TXT_FORMAT.")
    except Exception:
        # Geen aangeleverde tekst of persoonsgegevens in logs of foutmeldingen.
        return render_main(export_type=export_type, error="Verwerken niet gelukt. Foutcode: ERR_PROCESSING.")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
